"""第五期导入事务、导出一致性与不可变追问测试。"""

import io
import json
from pathlib import Path

import pytest
from docx import Document
from langgraph.checkpoint.memory import InMemorySaver
from sqlalchemy import create_engine, select
from sqlalchemy.orm import sessionmaker

from app.analysis_contracts import AnalysisRequest
from app.core.config import Settings
from app.db.index_models import IndexEvent
from app.db.models import Product
from app.db.report_models import ReportExport
from app.db.session import Base
from app.ingestion.imports import ImportRequest, import_records
from app.services.analysis_store import AnalysisStore
from app.services.reports import ask_report, create_export, load_report
from app.services.tasks import TaskService
from app.workflows.analysis import build_workflow


@pytest.fixture
def report_db():
    """隔离数据库导入一件商品，生成固定流程报告。"""
    engine = create_engine("sqlite://")
    Base.metadata.create_all(engine)
    factory = sessionmaker(engine, expire_on_commit=False)
    request = ImportRequest.model_validate({
        "products": [{"product_id": "p", "product_name": "密封盒", "market": "US",
                      "category": "home"}],
        "reviews": [{"review_id": "r", "product_id": "p", "market": "US", "category": "home",
                     "rating": 2, "variant": "大号", "content": "大号渗漏",
                     "observed_at": "2026-10-01T00:00:00Z", "data_kind": "synthetic"}],
    })
    imported = import_records(factory, request)
    settings = Settings(_env_file=None)
    store = AnalysisStore(factory, settings)
    payload = json.loads(Path("data/sample/analysis-request.json").read_text(encoding="utf-8"))
    payload["product_id"] = "p"
    identity = store.create(AnalysisRequest.model_validate(payload))
    build_workflow(store, identity, InMemorySaver()).invoke(
        {"run_id": identity}, {"configurable": {"thread_id": identity}},
    )
    yield factory, settings, identity, request, imported
    engine.dispose()


def test_import_repeat_conflict_and_atomic_rollback(report_db):
    factory, _, _, request, imported = report_db
    assert import_records(factory, request)["batch_id"] == imported["batch_id"]
    payload = request.model_dump(mode="json")
    payload["products"][0]["product_name"] = "不能覆盖的内容"
    payload["products"].insert(0, {"product_id": "new", "product_name": "先插入",
                                  "market": "US", "category": "home"})
    with pytest.raises(ValueError, match="不会覆盖"):
        import_records(factory, ImportRequest.model_validate(payload))
    with factory() as session:
        assert session.get(Product, "new") is None
        assert session.get(Product, "p").name == "密封盒"
        assert len(session.scalars(select(IndexEvent)).all()) == 1


def test_docx_numbers_citations_and_export_reuse(report_db):
    factory, _, identity, _, _ = report_db
    first = create_export(factory, identity)
    assert create_export(factory, identity) == first
    with factory() as session:
        row = session.get(ReportExport, first)
        doc = Document(io.BytesIO(row.content))
        text = "\n".join(p.text for p in doc.paragraphs)
        cells = [cell.text for table in doc.tables for row in table.rows for cell in row.cells]
        assert {"5.15", "7.25", "9.35"} <= set(cells)
        assert "ev-review-r" in text and "大号渗漏" in text
        assert len(session.scalars(select(ReportExport)).all()) == 1


def test_followup_keeps_original_report_hash(report_db):
    factory, settings, identity, _, _ = report_db
    original = load_report(factory, identity)
    response = ask_report(factory, settings, identity, "利润是多少")
    assert "7.25" in response["answer"]
    assert response["report_hash"] == original["report_hash"]
    assert load_report(factory, identity) == original


def test_cost_revision_preserves_snapshot_and_records_parent(report_db):
    factory, settings, identity, _, _ = report_db
    store = AnalysisStore(factory, settings)
    parent = store.get(identity)
    request = AnalysisRequest.model_validate(parent["request"])
    request.costs.sale_price = request.costs.sale_price + 1
    tasks = TaskService(factory, settings)
    child = tasks.create(request, "revision", snapshot_override=parent["snapshot"],
                         parent_id=identity)
    assert child != identity
    assert store.get(child)["snapshot"] == parent["snapshot"]
    assert store.get(child)["versions"]["parent_run_id"] == identity
    assert store.report(identity)["economics"]["scenarios"][1]["contribution"] == "7.25"
