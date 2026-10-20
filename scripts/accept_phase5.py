"""第五期真实数据库导入、三类报告、导出和追问验收。"""

import argparse
import io
import json
import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

from docx import Document  # noqa: E402

from app.analysis_contracts import AnalysisRequest  # noqa: E402
from app.core.config import get_settings  # noqa: E402
from app.db.report_models import ReportExport  # noqa: E402
from app.db.session import get_session_factory  # noqa: E402
from app.ingestion.imports import ImportRequest, import_records  # noqa: E402
from app.services.analysis_store import AnalysisStore  # noqa: E402
from app.services.indexing import sync_events  # noqa: E402
from app.services.reports import ask_report, create_export, load_report  # noqa: E402
from app.workflows.analysis import execute_analysis, postgres_checkpointer  # noqa: E402


def main():
    """--live 使冲突报告与一次追问使用真实模型，其他场景固定执行。"""
    parser = argparse.ArgumentParser()
    parser.add_argument("--live", action="store_true")
    args = parser.parse_args()
    settings, factory = get_settings(), get_session_factory()
    store = AnalysisStore(factory, settings)
    request = ImportRequest.model_validate_json(Path("data/workbench/import.json").read_text(
        encoding="utf-8",
    ))
    batch = import_records(factory, request)
    assert import_records(factory, request)["batch_id"] == batch["batch_id"]
    if args.live:
        indexed = sync_events(factory, settings)
        assert indexed["failed"] == 0, indexed
    template = json.loads(Path("data/sample/analysis-request.json").read_text(encoding="utf-8"))
    results = {"live": args.live, "import": batch, "reports": []}
    directory = Path("docs/examples")
    directory.mkdir(exist_ok=True)
    with postgres_checkpointer(settings) as saver:
        for name in ("normal", "missing", "conflict"):
            payload = {**template, "product_id": "demo-v5-" + name}
            if name == "missing":
                payload["costs"] = None
            if name == "conflict":
                payload["mode"] = "multi" if args.live else "fixed"
                payload["question"] = (
                    "分别委派 reviews 和 rules 核验500ml与1L的密封表现及使用条件。"
                    "明确是否可证明规格导致差异。摘要150字内，最多3条声明，记住是合成数据。"
                )
            identity = store.create(AnalysisRequest.model_validate(payload))
            print(f"{name} run_id={identity}", flush=True)
            report = execute_analysis(store, identity, saver)
            source = load_report(factory, identity)
            for format in ("json", "md", "docx"):
                export_id = create_export(factory, identity, format)
                assert create_export(factory, identity, format) == export_id
                with factory() as session:
                    row = session.get(ReportExport, export_id)
                    (directory / f"{name}.{format}").write_bytes(row.content)
                    if format == "docx":
                        document = Document(io.BytesIO(row.content))
                        text = "\n".join(p.text for p in document.paragraphs)
                        cells = [cell.text for table in document.tables
                                 for row in table.rows for cell in row.cells]
                        for claim in report["research"]["claims"]:
                            assert all(eid in text for eid in claim["evidence_ids"])
                        for scenario in report["economics"]["scenarios"]:
                            assert scenario["contribution"] in cells
            answer = ask_report(factory, settings, identity, "利润情况是什么")
            assert answer["report_hash"] == source["report_hash"]
            assert load_report(factory, identity) == source
            results["reports"].append({"case": name, "run_id": identity,
                                        "report_hash": source["report_hash"], "passed": True,
                                        "mode": report["mode"], "usage": report["usage"]})
            Path("docs/acceptance/phase-5-results.json").write_text(
                json.dumps(results, ensure_ascii=False, indent=2), encoding="utf-8",
            )
        if args.live:
            answer = ask_report(factory, settings, identity, "能否证明大号比小号更容易漏水？",
                                mode="model")
            assert load_report(factory, identity) == source
            results["live_followup"] = {"passed": True, "id": answer["id"],
                                        "answer": answer["answer"]}
    Path("docs/acceptance/phase-5-results.json").write_text(
        json.dumps(results, ensure_ascii=False, indent=2), encoding="utf-8",
    )
    print("phase5 passed", flush=True)


if __name__ == "__main__":
    main()
